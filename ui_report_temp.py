# ui_report_temp.py
import datetime
from PyQt5.QtWidgets import (QDialog, QVBoxLayout, QHBoxLayout, QTabWidget, 
                             QWidget, QLabel, QPushButton, QTableWidget, QTableWidgetItem,
                             QHeaderView, QComboBox, QMessageBox, QRadioButton, QFileDialog,
                             QListWidget, QAbstractItemView)
from PyQt5.QtCore import Qt
from PyQt5 import QtGui
import db_manager

# 엑셀 출력을 위한 openpyxl
try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
except ImportError:
    openpyxl = None

# matplotlib 및 pandas 임포트
import pandas as pd
import matplotlib.pyplot as plt
from matplotlib.backends.backend_qt5agg import FigureCanvasQTAgg as FigureCanvas
from matplotlib.figure import Figure
import matplotlib.ticker as ticker

# 한글 폰트 및 마이너스 부호 깨짐 방지
plt.rcParams['font.family'] = 'Malgun Gothic'
plt.rcParams['axes.unicode_minus'] = False


class TempReportDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("🌡 온도 및 부하 최고/최저 상관관계 분석 보고서")
        self.resize(1450, 850) # 열이 늘어나서 창 크기를 약간 더 넓혔습니다.
        self.setWindowFlags(self.windowFlags() | Qt.WindowMaximizeButtonHint | Qt.WindowMinimizeButtonHint)
        
        self.current_rows = []
        self.init_ui()

    def init_ui(self):
        main_layout = QVBoxLayout(self)
        
        title = QLabel("변압기 부하(kW) 및 실내/외/변압기 운전 온도(℃) 최고·최저 분석")
        title.setStyleSheet("font-size: 20px; font-weight: bold; color: #2c3e50;")
        title.setAlignment(Qt.AlignCenter)
        main_layout.addWidget(title)

        self.tabs = QTabWidget()
        self.tabs.setStyleSheet("QTabBar::tab { font-size: 14px; font-weight: bold; padding: 10px 20px; }")
        
        self.tab_table = QWidget()
        self.init_table_tab()
        self.tabs.addTab(self.tab_table, "📋 온도 및 부하 통계 데이터")
        
        self.tab_graph = QWidget()
        self.init_graph_tab()
        self.tabs.addTab(self.tab_graph, "📈 온도 추이 맞춤형 시각화")

        main_layout.addWidget(self.tabs)

        btn_layout = QHBoxLayout()
        btn_layout.addStretch()
        
        self.btn_export = QPushButton("📊 이 통계를 엑셀로 내보내기")
        self.btn_export.setStyleSheet("background-color: #27ae60; color: white; font-weight: bold; padding: 8px 15px;")
        self.btn_export.clicked.connect(self.export_to_excel)
        btn_layout.addWidget(self.btn_export)
        
        btn_close = QPushButton("닫기")
        btn_close.setStyleSheet("background-color: #7f8c8d; color: white; font-weight: bold; padding: 8px 30px;")
        btn_close.clicked.connect(self.accept)
        btn_layout.addWidget(btn_close)
        
        main_layout.addLayout(btn_layout)
        
        self.load_data()

    def init_table_tab(self):
        layout = QVBoxLayout(self.tab_table)
        ctrl_layout = QHBoxLayout()
        
        self.radio_month = QRadioButton("월간 조회 (일별 최고/최저)")
        self.radio_year = QRadioButton("연간 조회 (월별 최고/최저)")
        self.radio_month.setChecked(True)
        self.radio_month.toggled.connect(self.toggle_search_ui)
        
        ctrl_layout.addWidget(self.radio_month)
        ctrl_layout.addWidget(self.radio_year)
        ctrl_layout.addSpacing(20)
        
        self.combo_month = QComboBox()
        self.combo_year = QComboBox()
        
        now = datetime.datetime.now()
        for i in range(12):
            self.combo_month.addItem((now - datetime.timedelta(days=30*i)).strftime("%Y-%m"))
        for i in range(5): 
            self.combo_year.addItem(str(now.year - i))
            
        self.combo_year.setVisible(False) 
        
        ctrl_layout.addWidget(self.combo_month)
        ctrl_layout.addWidget(self.combo_year)
        
        btn_search = QPushButton("조회")
        btn_search.clicked.connect(self.load_data)
        ctrl_layout.addWidget(btn_search)
        ctrl_layout.addStretch()
        layout.addLayout(ctrl_layout)

        self.table = QTableWidget()
        # 💡 [핵심 보완] 총전력(KEP_P_kW) 항목 추가 (총 18열 데이터)
        self.data_labels = [
            "외기 최저(℃)", "외기 최고(℃)", 
            "실내 최저(℃)", "실내 최고(℃)", 
            "총전력 최저(kW)", "총전력 최고(kW)", 
            "TR1부하 최저", "TR1부하 최고", "TR1온도 최저", "TR1온도 최고", 
            "TR2부하 최저", "TR2부하 최고", "TR2온도 최저", "TR2온도 최고", 
            "TR3부하 최저", "TR3부하 최고", "TR3온도 최저", "TR3온도 최고"
        ]
        
        headers = ["구분(날짜/월)"] + self.data_labels
        self.table.setColumnCount(len(headers))
        self.table.setHorizontalHeaderLabels(headers)
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeToContents)
        
        # 안내 문구 디테일 업데이트
        warning_lbl = QLabel("※ 주의: 실내 35℃, TR온도 60℃ / 부하(총전력 1750kW, TR1 500kW, TR2·3 625kW) 이상 도달 시 적색 경고 표시")
        warning_lbl.setStyleSheet("color: #c0392b; font-weight: bold;")
        layout.addWidget(warning_lbl)
        layout.addWidget(self.table)

    def toggle_search_ui(self):
        if self.radio_month.isChecked():
            self.combo_month.setVisible(True)
            self.combo_year.setVisible(False)
        else:
            self.combo_month.setVisible(False)
            self.combo_year.setVisible(True)

    def init_graph_tab(self):
        layout = QVBoxLayout(self.tab_graph)
        
        graph_ctrl = QHBoxLayout()
        
        self.data_selector = QListWidget()
        self.data_selector.setSelectionMode(QAbstractItemView.MultiSelection) 
        self.data_selector.addItems(self.data_labels)
        self.data_selector.setMaximumHeight(100) 
        
        self.right_axis_selector = QListWidget()
        self.right_axis_selector.setSelectionMode(QAbstractItemView.MultiSelection)
        self.right_axis_selector.setMaximumHeight(100) 
        
        graph_ctrl.addWidget(QLabel("<b>[1] 그래프로 그릴 항목 선택:</b><br>(기본 왼쪽 축)"))
        graph_ctrl.addWidget(self.data_selector)
        graph_ctrl.addWidget(QLabel("<b>[2] 오른쪽 보조축(이중축)으로<br>보낼 항목 선택:</b>")) 
        graph_ctrl.addWidget(self.right_axis_selector)
        
        layout.addLayout(graph_ctrl)

        self.canvas = FigureCanvas(Figure(figsize=(10, 5)))
        self.ax = self.canvas.figure.add_subplot(111)
        layout.addWidget(self.canvas)

        self.data_selector.itemSelectionChanged.connect(self.sync_right_axis_list)
        self.right_axis_selector.itemSelectionChanged.connect(self.trigger_graph_update)

        # 초기 기본 선택 세팅 (총전력, 실내최고, TR1온도)
        self.data_selector.blockSignals(True)
        self.data_selector.item(1).setSelected(True) # 외기 최고
        self.data_selector.item(3).setSelected(True) # 실내 최고
        self.data_selector.item(5).setSelected(True) # 총전력 최고
        self.data_selector.item(9).setSelected(True) # TR1 온도 최고
        self.data_selector.blockSignals(False)
        self.sync_right_axis_list()

    def sync_right_axis_list(self):
        self.data_selector.blockSignals(True)
        self.right_axis_selector.blockSignals(True)
        try:
            prev_selected = [item.text() for item in self.right_axis_selector.selectedItems()]
            left_selected = [item.text() for item in self.data_selector.selectedItems()]
            
            self.right_axis_selector.clear()
            if left_selected:
                self.right_axis_selector.addItems(left_selected)
                for i in range(self.right_axis_selector.count()):
                    item = self.right_axis_selector.item(i)
                    if item.text() in prev_selected:
                        item.setSelected(True)
        finally:
            self.data_selector.blockSignals(False)
            self.right_axis_selector.blockSignals(False)
            
        self.trigger_graph_update()

    def trigger_graph_update(self):
        if hasattr(self, 'current_rows') and self.current_rows:
            self.update_graph(self.current_rows)

    def load_data(self):
        is_monthly = self.radio_month.isChecked()
        try:
            conn = db_manager.get_db_raw_connection()
            c = conn.cursor()
            
            # 💡 [핵심 보완] 총전력(KEP_P_kW) 컬럼을 쿼리에 추가
            if is_monthly:
                target = self.combo_month.currentText()
                query = """
                    SELECT DATE_FORMAT(log_date, '%%Y-%%m-%%d'), 
                           MIN(`외기온도`), MAX(`외기온도`), 
                           MIN(`실내온도`), MAX(`실내온도`),
                           MIN(`KEP_P_kW`), MAX(`KEP_P_kW`),
                           MIN(`Tr1_P_kW`), MAX(`Tr1_P_kW`), MIN(`Tr1_Temp`), MAX(`Tr1_Temp`), 
                           MIN(`Tr2_P_kW`), MAX(`Tr2_P_kW`), MIN(`Tr2_Temp`), MAX(`Tr2_Temp`), 
                           MIN(`Tr3_P_kW`), MAX(`Tr3_P_kW`), MIN(`Tr3_Temp`), MAX(`Tr3_Temp`)
                    FROM raw_data 
                    WHERE log_date LIKE %s 
                    GROUP BY log_date 
                    ORDER BY log_date ASC
                """
                c.execute(query, (f"{target}%",))
            else:
                target = self.combo_year.currentText()
                query = """
                    SELECT DATE_FORMAT(log_date, '%%Y-%%m'), 
                           MIN(`외기온도`), MAX(`외기온도`), 
                           MIN(`실내온도`), MAX(`실내온도`),
                           MIN(`KEP_P_kW`), MAX(`KEP_P_kW`),
                           MIN(`Tr1_P_kW`), MAX(`Tr1_P_kW`), MIN(`Tr1_Temp`), MAX(`Tr1_Temp`), 
                           MIN(`Tr2_P_kW`), MAX(`Tr2_P_kW`), MIN(`Tr2_Temp`), MAX(`Tr2_Temp`), 
                           MIN(`Tr3_P_kW`), MAX(`Tr3_P_kW`), MIN(`Tr3_Temp`), MAX(`Tr3_Temp`)
                    FROM raw_data 
                    WHERE log_date LIKE %s 
                    GROUP BY DATE_FORMAT(log_date, '%%Y-%%m') 
                    ORDER BY DATE_FORMAT(log_date, '%%Y-%%m') ASC
                """
                c.execute(query, (f"{target}%",))
                
            rows = c.fetchall()
            self.current_rows = rows 
            
            self.table.setRowCount(len(rows))
            for r_idx, row in enumerate(rows):
                for c_idx, val in enumerate(row):
                    item = QTableWidgetItem(f"{val:.1f}" if isinstance(val, float) else str(val))
                    item.setTextAlignment(Qt.AlignCenter)
                    
                    # 💡 [핵심 보완] 새로운 조건부 서식 (적색 경고 로직)
                    if isinstance(val, float):
                        # c_idx는 '구분(날짜)'가 0이므로 인덱스가 1씩 밀림
                        if c_idx == 4 and val >= 35.0: # 실내 최고
                            self.set_warning_format(item)
                        elif c_idx == 6 and val >= 1750.0: # 총전력 최고
                            self.set_warning_format(item)
                        elif c_idx == 8 and val >= 500.0: # TR1 부하 최고
                            self.set_warning_format(item)
                        elif c_idx in [12, 16] and val >= 625.0: # TR2, TR3 부하 최고
                            self.set_warning_format(item)
                        elif c_idx in [10, 14, 18] and val >= 60.0: # TR1, 2, 3 온도 최고 (60도로 수정)
                            self.set_warning_format(item)
                            
                    self.table.setItem(r_idx, c_idx, item)
            
            self.update_graph(rows)
            c.close(); conn.close()
            
        except Exception as e:
            print(f"온도 데이터 로딩 에러: {e}")
            QMessageBox.warning(self, "데이터 조회 오류", f"데이터베이스 조회 중 문제가 발생했습니다.\n에러: {e}")

    def set_warning_format(self, item):
        """적색 경고 서식을 일괄 적용하기 위한 헬퍼 함수"""
        item.setForeground(Qt.red)
        item.setFont(QtGui.QFont("Arial", 10, QtGui.QFont.Bold))

    def update_graph(self, rows):
        if not hasattr(self, 'canvas'): return
        
        self.canvas.figure.clf()
        self.ax = self.canvas.figure.add_subplot(111)
        
        if not rows:
            self.ax.text(0.5, 0.5, "조회된 데이터가 없습니다.", ha='center', va='center')
            self.canvas.draw()
            return
            
        selected_items = self.data_selector.selectedItems()
        if not selected_items:
            self.ax.text(0.5, 0.5, "위 리스트에서 그래프로 확인할 항목을 1개 이상 선택해 주세요.", ha='center', va='center')
            self.canvas.draw()
            return

        cols = ["날짜"] + self.data_labels
        df = pd.DataFrame(rows, columns=cols)
        
        df["날짜"] = df["날짜"].astype(str).apply(lambda x: x[5:] if len(x) == 10 else x)
        
        target_cols = [item.text() for item in selected_items]
        right_cols = [item.text() for item in self.right_axis_selector.selectedItems()]
        right_cols = [col for col in right_cols if col in target_cols] 
        left_cols = [col for col in target_cols if col not in right_cols]

        all_lines = []
        color_cycle = plt.rcParams['axes.prop_cycle'].by_key()['color']
        color_idx = 0

        # [기본 축 - 왼쪽]
        for col in left_cols:
            if col in df.columns:
                current_color = color_cycle[color_idx % len(color_cycle)]
                line = self.ax.plot(df["날짜"], df[col], marker='o', markersize=4, 
                                    color=current_color, label=col)
                all_lines += line
                color_idx += 1
        
        if left_cols:
            self.ax.set_ylabel("기본 축 (온도/부하)", color='#2c3e50', fontweight='bold', fontsize=11)
            self.ax.tick_params(axis='y', labelcolor='#2c3e50')
        else:
            self.ax.yaxis.set_visible(False)

        # [보조 축 - 오른쪽]
        if right_cols:
            ax2 = self.ax.twinx()
            for col in right_cols:
                if col in df.columns:
                    current_color = color_cycle[color_idx % len(color_cycle)]
                    line = ax2.plot(df["날짜"], df[col], marker='^', markersize=5, linestyle='--', 
                                    color=current_color, label=f"{col} (우측축)")
                    all_lines += line
                    color_idx += 1 
            
            ax2.set_ylabel("비교용 보조 축 (부하/온도)", color='#c0392b', fontweight='bold', fontsize=11)
            ax2.tick_params(axis='y', labelcolor='#c0392b')
            ax2.grid(False)

        if all_lines:
            labels = [l.get_label() for l in all_lines]
            if 'ax2' in locals():
                leg = ax2.legend(all_lines, labels, loc='upper left', bbox_to_anchor=(1.05, 1))
            else:
                leg = self.ax.legend(all_lines, labels, loc='upper left', bbox_to_anchor=(1.05, 1))
            
            self.canvas.figure.tight_layout()
            
        # 💡 [핵심 보완] X축 눈금(Tick) 표시 최적화
        if len(df) <= 12:
            # 연간 데이터(12개 이하)일 때는 무조건 매월(1칸 간격) 다 표시하도록 강제 설정
            self.ax.xaxis.set_major_locator(ticker.MultipleLocator(1))
        else:
            # 일간 데이터(31개 내외)일 때는 겹치지 않도록 12~15개 내외로 적절히 건너뛰며 표시
            self.ax.xaxis.set_major_locator(ticker.MaxNLocator(nbins=15, integer=True))
            
        self.ax.grid(True, linestyle=':', alpha=0.6)
        self.canvas.figure.autofmt_xdate(rotation=45) 
        
        self.canvas.draw()


    # =====================================================================
    # 엑셀 출력 기능
    # =====================================================================
    def export_to_excel(self):
        if openpyxl is None:
            QMessageBox.critical(self, "라이브러리 누락", "openpyxl 라이브러리가 설치되어 있지 않습니다.\n명령 프롬프트에서 'pip install openpyxl'을 실행하세요.")
            return

        timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M")
        
        if self.radio_month.isChecked():
            search_target = self.combo_month.currentText()
            report_title = f"온도부하_운전통계_월간({search_target})"
        else:
            search_target = self.combo_year.currentText()
            report_title = f"온도부하_운전통계_연간({search_target}년)"

        default_filename = f"{report_title}_{timestamp}.xlsx"
        save_path, _ = QFileDialog.getSaveFileName(self, "엑셀 파일 저장 위치 선택", default_filename, "Excel Files (*.xlsx)")
        
        if not save_path: return

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "온도 부하 데이터 통계"

            ws.page_setup.paperSize = ws.PAPERSIZE_A4
            ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0 
            
            ws.page_margins.left = 0.25
            ws.page_margins.right = 0.25
            ws.page_margins.top = 0.75
            ws.page_margins.bottom = 0.75

            col_count = self.table.columnCount()
            last_col_letter = openpyxl.utils.get_column_letter(col_count)
            
            title_range = f'A1:{last_col_letter}1'
            ws.merge_cells(title_range)
            cell_title = ws['A1']
            cell_title.value = "래미안개포루체하임아파트 전기설비 운전 온도(최고/최저) 데이터 통계"
            cell_title.font = Font(size=18, bold=True)
            cell_title.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[1].height = 35

            ws.merge_cells(f'A2:F2')
            ws['A2'].value = f"■ 조회 기준: {search_target}"
            ws['A2'].font = Font(bold=True)
            ws['A2'].alignment = Alignment(horizontal="left", vertical="center")
            
            date_range = f'G2:{last_col_letter}2'
            ws.merge_cells(date_range)
            print_time = datetime.datetime.now().strftime("%Y년 %m월 %d일 %H:%M")
            ws['G2'].value = f"■ 출력 일시: {print_time}"
            ws['G2'].font = Font(bold=True)
            ws['G2'].alignment = Alignment(horizontal="right", vertical="center")
            ws.row_dimensions[2].height = 20

            ws.row_dimensions[3].height = 10

            header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
            header_font = Font(bold=True)
            thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), 
                                 top=Side(style='thin'), bottom=Side(style='thin'))

            headers = [self.table.horizontalHeaderItem(c).text() for c in range(col_count)]
            for col_idx, h_text in enumerate(headers, start=1):
                cell = ws.cell(row=4, column=col_idx, value=h_text)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = thin_border
                ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 11 
                
            ws.column_dimensions['A'].width = 15 

            row_count = self.table.rowCount()
            for r in range(row_count):
                current_row = r + 5
                for c in range(col_count):
                    item = self.table.item(r, c)
                    text = item.text() if item else ""
                    
                    try:
                        val = float(text.replace(',', ''))
                    except ValueError:
                        val = text
                        
                    cell = ws.cell(row=current_row, column=c+1, value=val)
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.border = thin_border
                    
                    # 엑셀에서도 UI와 동일한 조건으로 적색 경고 적용
                    if isinstance(val, float):
                        is_warning = False
                        if c == 4 and val >= 35.0: is_warning = True
                        elif c == 6 and val >= 1750.0: is_warning = True
                        elif c == 8 and val >= 500.0: is_warning = True
                        elif c in [12, 16] and val >= 625.0: is_warning = True
                        elif c in [10, 14, 18] and val >= 60.0: is_warning = True
                        
                        if is_warning:
                            cell.font = Font(color="FF0000", bold=True)

            wb.save(save_path)
            QMessageBox.information(self, "출력 완료", f"A4 인쇄용 온도(최고/최저) 통계 엑셀 파일이 저장되었습니다.\n\n저장 위치:\n{save_path}")

        except Exception as e:
            QMessageBox.critical(self, "출력 실패", f"엑셀 파일 생성 중 오류가 발생했습니다.\n\n에러 내용: {e}")