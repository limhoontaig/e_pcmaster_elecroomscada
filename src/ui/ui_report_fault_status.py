# src/ui/ui_report_fault_status.py
import datetime
from PyQt5.QtWidgets import (QDialog, QVBoxLayout, QHBoxLayout, QWidget, QLabel, 
                             QPushButton, QTableWidget, QTableWidgetItem, QHeaderView, 
                             QComboBox, QMessageBox, QRadioButton, QFileDialog)
from PyQt5.QtCore import Qt
from PyQt5 import QtGui

from shared import db_manager

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
except ImportError:
    openpyxl = None

# =====================================================================
# 🔍 1. 원자료 상세 보기 팝업창 (특정 기기 장애 행 더블클릭 시 실행)
# =====================================================================
class FaultDataDetailDialog(QDialog):
    def __init__(self, target_period, equipment_name, event_msg, parent=None):
        super().__init__(parent)
        self.target_period = target_period
        self.equipment_name = equipment_name
        self.event_msg = event_msg
        
        self.setWindowTitle(f"🔍 [{self.target_period}] 장애 상세 이력 - {self.event_msg}")
        self.resize(1000, 600)
        self.init_ui()
        self.load_detail_data()

    def init_ui(self):
        layout = QVBoxLayout(self)
        
        lbl_title = QLabel(f"[{self.target_period}] '{self.event_msg}' 상세 발생 및 조치 내역")
        lbl_title.setStyleSheet("font-size: 16px; font-weight: bold; color: #c0392b;")
        layout.addWidget(lbl_title)

        self.table = QTableWidget()
        headers = ["발생 시간", "설비/태그명", "장애 내용", "조치 내역", "해제 시간", "조치 소요(분)", "조치자"]
        self.table.setColumnCount(len(headers))
        self.table.setHorizontalHeaderLabels(headers)
        
        header = self.table.horizontalHeader()
        header.setSectionResizeMode(QHeaderView.Interactive)
        self.table.setColumnWidth(0, 150) 
        self.table.setColumnWidth(1, 120) 
        self.table.setColumnWidth(2, 200) 
        self.table.setColumnWidth(3, 200) 
        self.table.setColumnWidth(4, 150) 
        self.table.setColumnWidth(5, 90)  
        self.table.setColumnWidth(6, 80)  
        
        self.table.horizontalHeader().setStretchLastSection(True)
        layout.addWidget(self.table)

        btn_close = QPushButton("닫기")
        btn_close.clicked.connect(self.accept)
        layout.addWidget(btn_close)

    def load_detail_data(self):
        try:
            conn = db_manager.get_db_raw_connection()
            c = conn.cursor()
            
            like_pattern = f"{self.target_period}%"
            
            query = """
                SELECT 
                    DATE_FORMAT(occurred_at, '%%Y-%%m-%%d %%H:%%i:%%s'),
                    equipment_name,
                    event_msg,
                    IFNULL(action_taken, '미조치'),
                    IFNULL(DATE_FORMAT(cleared_at, '%%Y-%%m-%%d %%H:%%i:%%s'), '진행중'),
                    IFNULL(ROUND(duration_seconds / 60.0, 1), 0.0),
                    operator_name
                FROM alarm_event_logs 
                WHERE DATE_FORMAT(occurred_at, '%%Y-%%m-%%d') LIKE %s
                  AND equipment_name = %s
                  AND event_type = 'ALARM'
                ORDER BY occurred_at DESC
            """
            c.execute(query, (like_pattern, self.equipment_name))
            rows = c.fetchall()
            
            self.table.setRowCount(len(rows))
            for r_idx, row in enumerate(rows):
                for c_idx, val in enumerate(row):
                    txt = str(val) if val is not None else "-"
                    item = QTableWidgetItem(txt)
                    item.setTextAlignment(Qt.AlignCenter)
                    
                    if c_idx == 3 and txt == '미조치':
                        item.setForeground(Qt.red)
                        item.setFont(QtGui.QFont("Arial", 10, QtGui.QFont.Bold))
                        
                    self.table.setItem(r_idx, c_idx, item)
                    
            c.close(); conn.close()
        except Exception as e:
            print(f"상세 데이터 로딩 에러: {e}")

# =====================================================================
# 🚨 2. 미조치 알람 전용 팝업창 
# =====================================================================
class PendingFaultsDialog(QDialog):
    def __init__(self, target_period, parent=None):
        super().__init__(parent)
        self.target_period = target_period
        
        self.setWindowTitle(f"🚨 [{self.target_period}] 현재 미조치 장애 상세 내역")
        self.resize(900, 500)
        self.init_ui()
        self.load_pending_data()

    def init_ui(self):
        layout = QVBoxLayout(self)
        
        lbl_title = QLabel(f"[{self.target_period}] 조치 대기 중인(미조치) 장애 전체 목록")
        lbl_title.setStyleSheet("font-size: 16px; font-weight: bold; color: #e74c3c;")
        layout.addWidget(lbl_title)

        self.table = QTableWidget()
        headers = ["발생 시간", "설비 태그명", "장애 내용", "조치 상태"]
        self.table.setColumnCount(len(headers))
        self.table.setHorizontalHeaderLabels(headers)
        
        header = self.table.horizontalHeader()
        header.setSectionResizeMode(QHeaderView.Stretch)
        layout.addWidget(self.table)

        btn_close = QPushButton("닫기")
        btn_close.clicked.connect(self.accept)
        layout.addWidget(btn_close)

    def load_pending_data(self):
        try:
            conn = db_manager.get_db_raw_connection()
            c = conn.cursor()
            
            like_pattern = f"{self.target_period}%"
            
            query = """
                SELECT 
                    DATE_FORMAT(occurred_at, '%%Y-%%m-%%d %%H:%%i:%%s'),
                    equipment_name,
                    event_msg,
                    IFNULL(action_taken, '미조치')
                FROM alarm_event_logs 
                WHERE DATE_FORMAT(occurred_at, '%%Y-%%m-%%d') LIKE %s
                  AND event_type = 'ALARM'
                  AND (cleared_at IS NULL OR action_taken = '미조치')
                ORDER BY occurred_at DESC
            """
            c.execute(query, (like_pattern,))
            rows = c.fetchall()
            
            self.table.setRowCount(len(rows))
            for r_idx, row in enumerate(rows):
                for c_idx, val in enumerate(row):
                    txt = str(val) if val is not None else "-"
                    item = QTableWidgetItem(txt)
                    item.setTextAlignment(Qt.AlignCenter)
                    
                    if c_idx == 3:
                        item.setForeground(Qt.red)
                        item.setFont(QtGui.QFont("Arial", 10, QtGui.QFont.Bold))
                        
                    self.table.setItem(r_idx, c_idx, item)
                    
            c.close(); conn.close()
        except Exception as e:
            print(f"미조치 데이터 로딩 에러: {e}")

# =====================================================================
# 📊 3. 메인 장애/알람 통계 보고서 창
# =====================================================================
class FaultStatusReportDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("🚨 설비 장애 및 알람 발생/조치 통계 보고서")
        self.resize(1200, 700)
        self.setWindowFlags(self.windowFlags() | Qt.WindowMaximizeButtonHint | Qt.WindowMinimizeButtonHint)
        self.init_ui()

    def init_ui(self):
        main_layout = QVBoxLayout(self)

        # 1. 타이틀
        title = QLabel("설비 장애(ALARM/TRIP) 발생 빈도 및 평균 조치 시간(MTTR) 분석")
        title.setStyleSheet("font-size: 20px; font-weight: bold; color: #c0392b;")
        title.setAlignment(Qt.AlignCenter)
        main_layout.addWidget(title)

        # 2. 조회 컨트롤 영역
        ctrl_layout = QHBoxLayout()
        
        self.radio_monthly = QRadioButton("월간 통계")
        self.radio_yearly = QRadioButton("연간 통계")
        self.radio_monthly.setChecked(True)

        self.radio_monthly.toggled.connect(self.toggle_ui)
        self.radio_yearly.toggled.connect(self.toggle_ui)

        ctrl_layout.addWidget(self.radio_monthly)
        ctrl_layout.addWidget(self.radio_yearly)
        ctrl_layout.addSpacing(20)

        self.combo_month = QComboBox()
        self.combo_year = QComboBox()
        
        now = datetime.datetime.now()
        for i in range(12):
            self.combo_month.addItem((now - datetime.timedelta(days=30*i)).strftime("%Y-%m"))
        for i in range(5):
            self.combo_year.addItem(str(now.year - i))
            
        self.combo_year.setVisible(False)
        
        ctrl_layout.addWidget(self.combo_month)
        ctrl_layout.addWidget(self.combo_year)

        btn_search = QPushButton("통계 조회")
        btn_search.setStyleSheet("background-color: #34495e; color: white; font-weight: bold; padding: 5px 15px;")
        btn_search.clicked.connect(self.load_statistics_data)
        ctrl_layout.addWidget(btn_search)
        
        ctrl_layout.addStretch()
        main_layout.addLayout(ctrl_layout)

        # 3. 요약 대시보드 영역 (원래 라벨 형태로 복구)
        summary_layout = QHBoxLayout()
        
        self.lbl_total_alarms = self.create_summary_card("총 장애 발생", "0 건", "#34495e")
        self.lbl_pending_alarms = self.create_summary_card("현재 미조치", "0 건", "#e74c3c")
        self.lbl_avg_repair_time = self.create_summary_card("평균 조치 시간", "0 분", "#27ae60")
        
        summary_layout.addWidget(self.lbl_total_alarms)
        summary_layout.addWidget(self.lbl_pending_alarms)
        summary_layout.addWidget(self.lbl_avg_repair_time)
        main_layout.addLayout(summary_layout)

        # 4. 통계 테이블
        self.table = QTableWidget()
        self.headers = ["설비/태그명(Hidden)", "발생 순위", "장애 내용", "설비 태그명", "발생 횟수", "미조치 건수", "평균 조치 소요시간(분)"]
        self.table.setColumnCount(len(self.headers))
        self.table.setHorizontalHeaderLabels(self.headers)
        
        self.table.setColumnHidden(0, True) 
        
        header = self.table.horizontalHeader()
        header.setSectionResizeMode(QHeaderView.Stretch)
        header.setSectionResizeMode(1, QHeaderView.ResizeToContents) 
        header.setSectionResizeMode(4, QHeaderView.ResizeToContents) 
        header.setSectionResizeMode(5, QHeaderView.ResizeToContents) 
        
        self.table.setAlternatingRowColors(True)
        self.table.cellDoubleClicked.connect(self.on_row_double_clicked)
        
        main_layout.addWidget(QLabel("※ 통계표의 항목을 더블클릭하면 해당 설비의 장애 발생 일시와 조치 내역 원본을 확인할 수 있습니다."))
        main_layout.addWidget(self.table)

        # 5. 하단 버튼 영역 (엑셀 버튼 좌측에 미조치 상세내역 버튼 추가)
        bottom_layout = QHBoxLayout()
        bottom_layout.addStretch()

        self.btn_show_pending = QPushButton("🚨 미조치 상세내역")
        self.btn_show_pending.setStyleSheet("background-color: #e74c3c; color: white; font-weight: bold; padding: 8px 15px;")
        self.btn_show_pending.clicked.connect(self.open_pending_dialog)
        bottom_layout.addWidget(self.btn_show_pending)

        self.btn_export = QPushButton("📊 엑셀 내보내기")
        self.btn_export.setStyleSheet("background-color: #27ae60; color: white; font-weight: bold; padding: 8px 15px;")
        self.btn_export.clicked.connect(self.export_to_excel)
        bottom_layout.addWidget(self.btn_export)

        btn_close = QPushButton("닫기")
        btn_close.setStyleSheet("background-color: #7f8c8d; color: white; font-weight: bold; padding: 8px 30px;")
        btn_close.clicked.connect(self.accept)
        bottom_layout.addWidget(btn_close)

        main_layout.addLayout(bottom_layout)

        self.load_statistics_data()

    def create_summary_card(self, title, value, color):
        lbl = QLabel(f"{title}\n{value}")
        lbl.setAlignment(Qt.AlignCenter)
        lbl.setStyleSheet(f"""
            background-color: {color}; color: white; 
            font-size: 16px; font-weight: bold; 
            padding: 15px; border-radius: 8px;
        """)
        return lbl

    def update_summary_card(self, lbl_widget, title, value):
        lbl_widget.setText(f"{title}\n{value}")

    def toggle_ui(self):
        if self.radio_monthly.isChecked():
            self.combo_month.setVisible(True)
            self.combo_year.setVisible(False)
        else:
            self.combo_month.setVisible(False)
            self.combo_year.setVisible(True)
        self.load_statistics_data()

    def load_statistics_data(self):
        if self.radio_monthly.isChecked():
            target_period = self.combo_month.currentText()
        else:
            target_period = self.combo_year.currentText()
            
        like_pattern = f"{target_period}%"

        try:
            conn = db_manager.get_db_raw_connection()
            c = conn.cursor()

            query = """
                SELECT 
                    equipment_name,
                    event_msg,
                    COUNT(event_id) as freq_count,
                    SUM(CASE WHEN cleared_at IS NULL OR action_taken = '미조치' THEN 1 ELSE 0 END) as pending_count,
                    IFNULL(AVG(duration_seconds), 0) as avg_duration
                FROM alarm_event_logs
                WHERE DATE_FORMAT(occurred_at, '%%Y-%%m-%%d') LIKE %s
                  AND event_type = 'ALARM'
                GROUP BY equipment_name, event_msg
                ORDER BY freq_count DESC
            """
            c.execute(query, (like_pattern,))
            rows = c.fetchall()
            
            total_alarms = 0
            total_pending = 0
            sum_duration = 0.0
            
            self.table.setRowCount(len(rows))
            
            for r_idx, row in enumerate(rows):
                equip_name, event_msg, freq_count, pending_count, avg_duration_sec = row
                
                total_alarms += freq_count
                total_pending += int(pending_count)
                sum_duration += (float(avg_duration_sec) * freq_count) 
                
                avg_min = round(float(avg_duration_sec) / 60.0, 1)
                
                self.set_table_item(r_idx, 0, equip_name)
                self.set_table_item(r_idx, 1, f"{r_idx + 1}위")
                self.set_table_item(r_idx, 2, event_msg)
                self.set_table_item(r_idx, 3, equip_name)
                self.set_table_item(r_idx, 4, f"{freq_count}회")
                
                item_pending = QTableWidgetItem(f"{pending_count}건")
                item_pending.setTextAlignment(Qt.AlignCenter)
                if int(pending_count) > 0:
                    item_pending.setForeground(Qt.red)
                    item_pending.setFont(QtGui.QFont("Arial", 10, QtGui.QFont.Bold))
                self.table.setItem(r_idx, 5, item_pending)
                
                self.set_table_item(r_idx, 6, f"{avg_min}분")

            c.close(); conn.close()
            
            # 요약 대시보드 갱신
            self.update_summary_card(self.lbl_total_alarms, "총 장애 발생", f"{total_alarms} 건")
            self.update_summary_card(self.lbl_pending_alarms, "현재 미조치", f"{total_pending} 건")
            
            overall_avg_min = round((sum_duration / total_alarms) / 60.0, 1) if total_alarms > 0 else 0.0
            self.update_summary_card(self.lbl_avg_repair_time, "평균 조치 시간", f"{overall_avg_min} 분")

        except Exception as e:
            print(f"장애 통계 로딩 에러: {e}")
            QMessageBox.warning(self, "데이터 조회 오류", f"통계 데이터를 불러오는 중 오류가 발생했습니다.\n{e}")

    def set_table_item(self, row, col, text):
        item = QTableWidgetItem(text)
        item.setTextAlignment(Qt.AlignCenter)
        self.table.setItem(row, col, item)

    def on_row_double_clicked(self, row, column):
        target_period = self.combo_month.currentText() if self.radio_monthly.isChecked() else self.combo_year.currentText()
        equip_name = self.table.item(row, 0).text()
        event_msg = self.table.item(row, 2).text()
        
        dialog = FaultDataDetailDialog(target_period, equip_name, event_msg, self)
        dialog.exec_()
        
    def open_pending_dialog(self):
        """🌟 하단 버튼 클릭 시 미조치 알람 팝업 실행"""
        target_period = self.combo_month.currentText() if self.radio_monthly.isChecked() else self.combo_year.currentText()
        
        dialog = PendingFaultsDialog(target_period, self)
        dialog.exec_()

    def export_to_excel(self):
        if openpyxl is None:
            QMessageBox.critical(self, "라이브러리 누락", "openpyxl 라이브러리가 설치되어 있지 않습니다.")
            return

        target_str = self.combo_month.currentText() if self.radio_monthly.isChecked() else self.combo_year.currentText()
        period = "월간" if self.radio_monthly.isChecked() else "연간"

        timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M")
        default_filename = f"장애발생_{period}통계_{target_str}_{timestamp}.xlsx"
        
        save_path, _ = QFileDialog.getSaveFileName(self, "엑셀 파일 저장 위치 선택", default_filename, "Excel Files (*.xlsx)")
        
        if not save_path:
            return

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = f"{period} 알람 통계"

            col_count = self.table.columnCount()
            headers = [self.table.horizontalHeaderItem(c).text() for c in range(1, col_count)]
            ws.append(headers)

            header_fill = PatternFill(start_color="FADBD8", end_color="FADBD8", fill_type="solid")
            thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), 
                                 top=Side(style='thin'), bottom=Side(style='thin'))

            for col_idx in range(1, col_count):
                cell = ws.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin_border
                ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 20

            ws.column_dimensions['B'].width = 35 

            row_count = self.table.rowCount()
            for r in range(row_count):
                row_data = []
                for c in range(1, col_count): 
                    item = self.table.item(r, c)
                    text = item.text().replace('회', '').replace('건', '').replace('위', '').replace('분', '').strip() if item else ""
                    
                    try:
                        val = float(text)
                        if val.is_integer(): val = int(val)
                        row_data.append(val)
                    except ValueError:
                        row_data.append(item.text() if item else "")
                
                ws.append(row_data)
                
                current_row = r + 2
                for col_idx in range(1, col_count):
                    cell = ws.cell(row=current_row, column=col_idx)
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.border = thin_border

            wb.save(save_path)
            QMessageBox.information(self, "출력 완료", f"장애 통계 데이터가 엑셀로 저장되었습니다.\n\n위치: {save_path}")

        except Exception as e:
            QMessageBox.critical(self, "출력 실패", f"엑셀 파일 생성 중 오류가 발생했습니다.\n{e}")